import ga_params
if ga_params.draw_plot:
    import matplotlib.pyplot as plt
if ga_params.export_spreadsheet:
    from openpyxl import Workbook
from datetime import datetime


class Reporter:
    
    #Stocke les génerations et résultats
    plot_x_axis = []
    plot_y_axis = []
    
    # workbook (classeur: fichier) et worksheet(feuille de calcul) Excel
    sheet_wb = None
    sheet_ws = None
    
    #Nom du fichier Excel
    sheet_dest = 'output'
    sheet_dest_ext = '.xlsm'

    def __init__(self, run_name: str):
        """self.fig, self.subplot, self.fig_node_connector, self.subplot2 et self.init_spreadsheet sont dans __init__ (attribut d'instance) parce que :
            * ils dépendent d’un paramètre (draw_plot ou export_spreadsheet)
            * ils sont créés dynamiquement
            * chaque instance doit avoir sa propre figure et son propre spreadsheet"""
        # Initialier le graphique
        if ga_params.draw_plot:
            plt.ion() # Mode interactif de matplotlib(mise à jour en temps réel des graphiques sans bloquer l'exécution de l'algorithme)
            # crée 2 figures
            self.fig, self.subplot = plt.subplots(figsize=(20, 10)) # Figure 1 → convergence + tableau 
            self.fig_node_connector, self.subplot2 = plt.subplots() # Figure 2 → carte des routes
            
        # Initialier le fichier Excel
        if ga_params.export_spreadsheet:
            self.init_spreadsheet(run_name=run_name)

    # fonction qui gènère les 2 figures (plot-output.png et plot2-output.png)
    def plot_draw(self, x_axis: list, y_axis: list, latest_result, total_time=None):
        # Extraire les valeurs float de chaque resultat
        y_values = []
        for y in y_axis:
            try:
                y_values.append(float(y['best'].value)) # Extrait la valeur du meilleur chromosome (best.value : la val min de la fn objective) de la generation et la convertit en float
            except:
                break # Si erreur (ex: valeur manquante), on sort de la boucle
        if not y_values:
            return
        
        # Formater le temps total
        if total_time:
            total_str = str(total_time).split('.')[0]  # enlever les microsecondes (Exemple: "0:15:23.456789" → "0:15:23" )
        else:
            total_str = 'N/A'
        
         # ── Figure 1 ────────────────────────────────────────────────────────────
        plt.figure(1)
        plt.clf()
        self.fig.set_size_inches(28, 14)   # figure plus large pour le tableau

        plt.suptitle(
            f'Gen {x_axis[-1]} | Best: {latest_result.value:.1f} | '
            f'Vehicles: {latest_result.vehicles_count} | Temps total: {total_str}',
            fontsize=11
        )

        # Courbe de convergence — partie haute (30 % de la hauteur)
        ax1 = plt.subplot2grid((10, 1), (0, 0), rowspan=3)
        margin = (max(y_values) - min(y_values)) * 0.1 + 10
        ax1.set_ylim([min(y_values) - margin, max(y_values) + margin])
        ax1.set_xlabel('Generation')
        ax1.set_ylabel('Cost')
        ax1.plot(x_axis, y_values, color='steelblue')

        # Tableau des routes — partie basse (70 % de la hauteur)
        ax2 = plt.subplot2grid((10, 1), (3, 0), rowspan=7)
        ax2.axis('off')

        # Construction des données
        col_labels = ['Véhicule', 'Client', 'Arrivée', 'Début TW', 'Fin TW', 'Départ']
        table_data = []

        for i, (route, timings) in enumerate(
                zip(latest_result.vehicles_routes, latest_result.vehicles_timings)):
            for t in timings:
                late = t['arrival'] > t['due_time']
                wait = t['arrival'] < t['ready_time']
                flag = ' [!]' if late else (' [w]' if wait else '')   # ASCII pur
                table_data.append([
                    f'V{i + 1}',
                    str(t['node']) + flag,
                    f"{t['arrival']:.1f}",
                    f"{t['ready_time']}",
                    f"{t['due_time']}",
                    f"{t['departure']:.1f}",
                ])

        if not table_data:
            plt.tight_layout()
            plt.draw()
            self.fig.savefig("plot-output.png")
            plt.pause(0.000001)
            return

        table = ax2.table(
            cellText=table_data,
            colLabels=col_labels,
            cellLoc='center',
            loc='center',
            bbox=[0.05, 0, 0.9, 1]   # occupe tout ax2
        )
        table.auto_set_font_size(False)
        table.set_fontsize(6.5)

        # Largeurs relatives : Véhicule, Client, Arrivée, Début TW, Fin TW, Départ
        col_widths = [0.001, 0.001, 0.001, 0.001, 0.001, 0.001]
        for (row, col), cell in table.get_celld().items():
            cell.set_width(col_widths[col] if col < len(col_widths) else 0.10)
            cell.set_linewidth(0.3)

        # En-tête
        for j in range(len(col_labels)):
            table[0, j].set_facecolor('#4472C4')
            table[0, j].set_text_props(color='white', fontweight='bold')

        # Lignes alternées + colorisation des violations
        for i in range(1, len(table_data) + 1):
            raw_flag = table_data[i - 1][1]          # colonne Client
            is_late  = '[!]' in raw_flag
            is_wait  = '[w]' in raw_flag
            base_color = '#FFD0D0' if is_late else ('#FFF3CC' if is_wait else
                        ('#DCE6F1' if i % 2 == 0 else 'white'))
            for j in range(len(col_labels)):
                table[i, j].set_facecolor(base_color)

        plt.tight_layout(rect=[0, 0, 1, 0.97])
        
        
        # Figure 2 : carte des routes
        plt.figure(2)
        plt.clf()
        # Parcourt chaque tournée
        # route_x : Liste des coordonnées X des points de la tournée (clients + depôt)
        # route_y : liste des coordonnées Y des points de la tournée (noeuds)
        for route_x, route_y in latest_result.plot_get_route_cords(): 
            plt.plot(route_x, route_y, marker='o', markersize=3) #  Trace la ligne reliant les points et marker='o' ajoute des cercles aux points
        plt.title(f'Routes | {latest_result.vehicles_count} vehicles') # Ajoute un titre avec le nombre de véhicules

        plt.tight_layout(pad=2.0) # Ajuste la mise en page avec un padding de 2 points
        plt.draw() # Force le rafraîchissement de l'affichage
        self.fig.savefig("plot-output.png")
        self.fig_node_connector.savefig("plot2-output.png")
        plt.pause(0.000001) # Fait une pause très courte (1 microseconde) pour permettre l'affichage (nécessaire en mode interactif)

    # Cree le fichier Excel nommé automatiquement avec la date 
    def init_spreadsheet(self, run_name: str):
        datetime_str = datetime.now().strftime('%b%d-%H%M%S')  # Format: Mar27-155523 (sans les deux-points)
        self.sheet_dest = self.sheet_dest + '_' + run_name + '_' + datetime_str + self.sheet_dest_ext
        self.sheet_wb = Workbook()
        self.sheet_ws = self.sheet_wb.active
        self.sheet_ws.append(['gen_index', 'best', 'average', 'std', 'worst', 'created_time', 'computation_time',
                              'best_route'])

    # Ecrit et sauvegarde le fichier Excel
    def export_spreadsheet(self, x_axis: list, y_axis: list):
        datetime_str = str(datetime.now())
        for x, y in zip(x_axis, y_axis):
            self.sheet_ws.append([x, # numéro de génération
                                  y['best'].value, # valeur objective du meilleur
                                  y['average'], # moyenne de la population
                                  y['std'], # écart-type (diversité)
                                  y['worst'].value, # valeur objective du pire
                                  datetime_str, # horodatage
                                  str(y['process_time']), # temps de calcul
                                  str(y['best'].route)]) # route du meilleur chromosome
        self.sheet_wb.save(filename=self.sheet_dest)

    # Export final unique avec rapport détaillé des résultats
    def export_spreadsheet_final(self, best_solution, total_time, generation_count):
        """
        Génère un rapport Excel final unique contenant:
        - Onglet "Résumé Final": Statistiques globales et meilleure solution
        - Onglet "Routes Détaillées": Détails de chaque route avec horaires
        Et génère les graphiques PNG à la fin
        """
        try:
            # Créer un workbook avec le résumé final
            self.sheet_wb = Workbook()
            self.sheet_ws = self.sheet_wb.active
            self.sheet_ws.title = "Résumé Final"
            
            # Ajouter les headers et données du résumé
            self.sheet_ws.append(['Métrique', 'Valeur'])
            self.sheet_ws.append(['Meilleur Coût (Fonction Objective)', round(best_solution.value, 2)])
            self.sheet_ws.append(['Temps d\'Exécution Total (sec)', round(total_time.total_seconds(), 2)])
            self.sheet_ws.append(['Nombre de Véhicules', best_solution.vehicles_count])
            self.sheet_ws.append(['Nombre de Générations', generation_count])
            self.sheet_ws.append(['Distance Totale Parcourue', round(best_solution.total_travel_dist, 2)])
            
            # Créer un nouvel onglet pour les routes détaillées
            ws_routes = self.sheet_wb.create_sheet("Routes Détaillées")
            ws_routes.append(['Véhicule', 'Client', 'Heure d\'Arrivée (sec)', 'Ready Time (sec)', 'Due Time (sec)', 
                             'Heure de Départ (sec)', 'Temps d\'Attente (sec)'])
            
            # Remplir les détails des routes
            for vehicle_idx, vehicle_timings in enumerate(best_solution.vehicles_timings):
                for timing in vehicle_timings:
                    # Calcul du temps d'attente réel
                    waiting_time = max(0, timing['ready_time'] - timing['arrival'])
                    ws_routes.append([
                        f'V{vehicle_idx + 1}',  # Véhicule
                        timing['node'],  # Client
                        round(timing['arrival'], 2),  # Heure d'arrivée
                        timing['ready_time'],  # Ready time (earliest service start)
                        timing['due_time'],  # Due time (latest service start)
                        round(timing['departure'], 2),  # Heure de départ
                        round(waiting_time, 2)  # Temps d'attente
                    ])
            
            # Sauvegarder le fichier
            self.sheet_wb.save(filename=self.sheet_dest)
            print(f"\n✓ Rapport Excel généré: {self.sheet_dest}")
            
            # Générer les graphiques à la fin
            if ga_params.draw_plot:
                self.plot_draw_final(best_solution=best_solution, total_time=total_time, generation_count=generation_count)
        
        except Exception as e:
            print(f"✗ Erreur lors de la génération du rapport: {type(e).__name__}: {e}")
            import traceback
            traceback.print_exc()
    
    # Génère les graphiques finaux
    def plot_draw_final(self, best_solution, total_time, generation_count):
        """Génère les graphiques PNG finaux"""
        if not ga_params.draw_plot:
            return
        
        import matplotlib.pyplot as plt
        
        # Figure 1 : Tableau des routes et informations résumées
        fig1 = plt.figure(figsize=(12, 8))
        ax1 = plt.subplot(111)
        ax1.axis('off')
        
        total_seconds = total_time.total_seconds()
        hours = int(total_seconds // 3600)
        minutes = int((total_seconds % 3600) // 60)
        seconds = int(total_seconds % 60)
        time_str = f'{hours:02d}:{minutes:02d}:{seconds:02d}'
        
        plt.suptitle(
            f'Solution Finale | Coût: {best_solution.value:.2f} | Véhicules: {best_solution.vehicles_count} | Temps: {time_str}',
            fontsize=12
        )
        
        # Créer le tableau des routes
        table_data = []
        for i, route in enumerate(best_solution.vehicles_routes):
            trajet = ' → '.join(str(n) for n in route)
            nb_clients = len(route) - 2
            table_data.append([f'V{i+1}', trajet, str(nb_clients)])
        
        table = ax1.table(
            cellText=table_data,
            colLabels=['Véhicule', 'Trajet', 'Nb Clients'],
            cellLoc='left',
            loc='center'
        )
        table.auto_set_font_size(False)
        table.set_fontsize(8)
        table.auto_set_column_width([0, 1, 2])
        
        for j in range(3):
            table[0, j].set_facecolor('#4472C4')
            table[0, j].set_text_props(color='white', fontweight='bold')
        
        plt.tight_layout(rect=[0, 0, 1, 0.97])
        fig1.savefig("plot-output.png")
        plt.close(fig1)
        
        # Figure 2 : Carte des routes
        fig2 = plt.figure(figsize=(12, 8))
        for route_x, route_y in best_solution.plot_get_route_cords(): 
            plt.plot(route_x, route_y, marker='o', markersize=3)
        plt.title(f'Routes Finales | {best_solution.vehicles_count} véhicules | Génération {generation_count}')
        plt.xlabel('X')
        plt.ylabel('Y')
        plt.grid(True, alpha=0.3)
        plt.tight_layout()
        fig2.savefig("plot2-output.png")
        plt.close(fig2)
        
